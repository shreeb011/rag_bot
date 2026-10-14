import io

import docx
import pytest

from app.errors import DocumentParseError, UnsupportedFileTypeError
from app.ingestion.chunker import chunk_sections, split_text
from app.ingestion.loaders import load_document
from app.models import Section
from tests.conftest import SAMPLE_DOCS


def test_markdown_sections_keep_heading_path():
    md = b"# Handbook\nIntro.\n## Leave\n### Annual\n24 days of leave.\n## Equipment\nLaptop issued.\n"
    sections = load_document("h.md", md)
    assert [(s.heading, s.text) for s in sections] == [
        ("Handbook", "Intro."),
        ("Handbook > Leave > Annual", "24 days of leave."),
        ("Handbook > Equipment", "Laptop issued."),
    ]


def test_markdown_ignores_hash_lines_inside_code_blocks():
    md = b"# Title\n```\n# not a heading\n```\ntext\n"
    sections = load_document("c.md", md)
    assert len(sections) == 1 and "# not a heading" in sections[0].text


def test_txt_decodes_utf8_bom_and_cp1252():
    assert load_document("a.txt", "\ufeffcaf\u00e9".encode("utf-8"))[0].text == "caf\u00e9"
    assert load_document("b.txt", "caf\u00e9".encode("cp1252"))[0].text == "caf\u00e9"


def test_pdf_extraction_keeps_page_numbers():
    sections = load_document("policy.pdf", (SAMPLE_DOCS / "information_security_policy.pdf").read_bytes())
    assert [s.page for s in sections] == [1, 2, 3]
    assert "at least 14 characters" in sections[0].text
    assert "within 4 hours" in sections[1].text


def _make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Handbook", level=0)
    document.add_paragraph("Welcome to Nimbus.")
    document.add_heading("Leave", level=1)
    document.add_heading("Annual", level=2)
    document.add_paragraph("Employees get 24 days of paid leave.")
    document.add_paragraph("Carry over up to 5 days", style="List Bullet")
    document.add_heading("Plans", level=1)
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Plan", "SLA"), ("Pro", "99.9%")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_keeps_heading_path_lists_and_tables():
    sections = load_document("handbook.docx", _make_docx())
    assert [(s.heading, s.text) for s in sections] == [
        ("Handbook", "Welcome to Nimbus."),
        ("Handbook > Leave > Annual", "Employees get 24 days of paid leave.\n\n- Carry over up to 5 days"),
        ("Handbook > Plans", "Plan | SLA\nPro | 99.9%"),
    ]


def _make_resume_docx() -> bytes:
    """A CV laid out the way people usually write them: no heading styles, just bold / ALL-CAPS lines."""
    document = docx.Document()
    for line in ["JANE DOE M.S.", "DATA ENGINEER"]:
        document.add_paragraph().add_run(line).bold = True
    document.add_paragraph("Email: jane@example.com | Phone: +1 555 0100 | Location: Austin, TX")
    document.add_paragraph().add_run("PROFESSIONAL SUMMARY").bold = True
    document.add_paragraph("Data engineer with 6 years of experience.")
    document.add_paragraph().add_run("Acme Corp - Senior Engineer 2021 - Present").bold = True
    document.add_paragraph("Built streaming pipelines.")
    document.add_paragraph().add_run("SKILLS").bold = True
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Languages:", "Python, SQL"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_without_heading_styles_uses_bold_and_caps_lines():
    sections = load_document("cv.docx", _make_resume_docx())
    assert [(s.heading, s.text) for s in sections] == [
        ("JANE DOE M.S.", "JANE DOE M.S.\nDATA ENGINEER\nEmail: jane@example.com\nPhone: +1 555 0100\nLocation: Austin, TX"),
        ("JANE DOE M.S. > PROFESSIONAL SUMMARY", "Data engineer with 6 years of experience."),
        ("JANE DOE M.S. > PROFESSIONAL SUMMARY > Acme Corp - Senior Engineer 2021 - Present", "Built streaming pipelines."),
        ("JANE DOE M.S. > SKILLS", "Languages: Python, SQL"),
    ]


@pytest.mark.parametrize(
    "line,expected",
    [
        ("Email: a@b.com | Phone: 123 \u2022 Location: X", "Email: a@b.com\nPhone: 123\nLocation: X"),
        ("| Plan | SLA |\n| Pro | 99.9% |", "| Plan | SLA |\n| Pro | 99.9% |"),
        ("See https://x.io | docs", "See https://x.io | docs"),
    ],
)
def test_inline_fields_split_but_tables_untouched(line, expected):
    assert load_document("f.txt", line.encode())[0].text == expected


def test_legacy_doc_rejected_with_hint():
    with pytest.raises(UnsupportedFileTypeError, match=r"\.docx"):
        load_document("old.doc", b"\xd0\xcf\x11\xe0")


@pytest.mark.parametrize("name", ["evil.exe", "slides.pptx", "noextension"])
def test_unsupported_extension_rejected(name):
    with pytest.raises(UnsupportedFileTypeError):
        load_document(name, b"hello")


@pytest.mark.parametrize(
    "name,data",
    [("empty.txt", b""), ("blank.md", b"   \n\n  "), ("broken.pdf", b"%PDF-1.4 garbage"), ("broken.docx", b"PK not a zip")],
)
def test_empty_or_corrupt_documents_rejected(name, data):
    with pytest.raises(DocumentParseError):
        load_document(name, data)


def test_chunks_respect_size_and_overlap():
    paragraphs = [f"Paragraph {i} " + "word " * 30 for i in range(20)]
    chunks = split_text("\n\n".join(paragraphs), chunk_size=400, chunk_overlap=200)
    assert len(chunks) > 1
    assert all(len(c) <= 400 for c in chunks)
    for previous, current in zip(chunks, chunks[1:]):
        assert previous.split("\n")[-1] == current.split("\n")[0], "consecutive chunks should share a paragraph"


def test_oversized_paragraph_is_split_on_sentences_and_words():
    text = "Sentence one is here. " * 40 + "x" * 1200
    chunks = split_text(text, chunk_size=300, chunk_overlap=50)
    assert all(len(c) <= 300 for c in chunks)
    assert "".join(chunks).count("x") >= 1200


def test_chunk_metadata_and_ids():
    sections = [Section("First page text.", page=1), Section("Second page text.", page=2, heading="H")]
    chunks = chunk_sections(sections, doc_id="d1", doc_name="doc.pdf", chunk_size=200, chunk_overlap=20)
    assert [(c.id, c.page, c.heading) for c in chunks] == [("d1:0", 1, None), ("d1:1", 2, "H")]
    assert chunks[1].embedding_text().startswith("doc.pdf > H\n")


def test_overlap_must_be_smaller_than_chunk_size():
    with pytest.raises(ValueError):
        chunk_sections([Section("x")], "d", "n", chunk_size=100, chunk_overlap=100)
