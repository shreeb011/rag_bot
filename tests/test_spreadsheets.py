"""Excel (.xlsx / .xls) and CSV ingestion: one section per sheet, one self-describing line per row."""

from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

import openpyxl
import pytest

from app.errors import DocumentParseError
from app.ingestion.loaders import load_document

FIXTURES = Path(__file__).parent / "fixtures"


def _xlsx(sheets: dict[str, list[list]], formats: dict[tuple[str, str], str] | None = None) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    for (title, ref), number_format in (formats or {}).items():
        workbook[title][ref].number_format = number_format
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_rows_repeat_column_names_and_keep_row_numbers():
    data = _xlsx(
        {
            "Employees": [
                ["Employee list 2025"],
                [],
                ["Name", "Department", "Salary", "Bonus", "Joined", "Active"],
                ["Alice Rao", "Sales", 52000, 0.15, dt.datetime(2021, 3, 1), True],
                ["Bob Iyer", "Engineering", 61000.5, None, dt.datetime(2019, 7, 15, 9, 30), False],
            ]
        },
        formats={("Employees", "D4"): "0%"},
    )
    (section,) = load_document("staff.xlsx", data)
    assert section.heading == "Employees"
    assert section.text.split("\n\n") == [
        "Employee list 2025",
        "Row 4 | Name: Alice Rao; Department: Sales; Salary: 52000; Bonus: 15%; Joined: 2021-03-01; Active: Yes",
        "Row 5 | Name: Bob Iyer; Department: Engineering; Salary: 61000.5; Joined: 2019-07-15 09:30; Active: No",
    ]


def test_xlsx_one_section_per_sheet_and_empty_sheets_skipped():
    data = _xlsx({"Offices": [["City", "Head count"], ["Chennai", 40]], "Blank": [], "Notes": [["Reviewed in Q3"]]})
    sections = load_document("company.xlsx", data)
    assert [(s.heading, s.text) for s in sections] == [
        ("Offices", "Row 2 | City: Chennai; Head count: 40"),
        ("Notes", "Reviewed in Q3"),
    ]


def test_sheet_without_header_row_keeps_plain_values():
    data = _xlsx({"Data": [[1, 2, 3], [4, 5, 6]]})
    (section,) = load_document("numbers.xlsx", data)
    assert section.text == "Row 1 | 1; 2; 3\n\nRow 2 | 4; 5; 6"


def test_blank_header_cells_get_column_letters():
    data = _xlsx({"S": [["Name", None, "Score"], ["Ann", "extra", 9]]})
    (section,) = load_document("x.xlsx", data)
    assert section.text == "Row 2 | Name: Ann; Column B: extra; Score: 9"


def test_xls_legacy_workbook():
    (section,) = load_document("inventory.xls", (FIXTURES / "inventory.xls").read_bytes())
    assert section.heading == "Inventory"
    assert section.text == "Row 2 | Item: Laptop; Qty: 12; Received: 2024-05-02\n\nRow 3 | Item: Monitor; Qty: 30; Received: 2024-06-09"


@pytest.mark.parametrize(
    "content",
    [
        "Product,Price,In stock\nWidget,9.99,yes\nGadget,19.50,no\n",
        "Product;Price;In stock\r\nWidget;9.99;yes\r\nGadget;19.50;no\r\n",
        "Product\tPrice\tIn stock\nWidget\t9.99\tyes\nGadget\t19.50\tno\n",
    ],
    ids=["comma", "semicolon", "tab"],
)
def test_csv_delimiters_are_detected(content):
    (section,) = load_document("products.csv", content.encode())
    assert section.heading is None
    assert section.text == "Row 2 | Product: Widget; Price: 9.99; In stock: yes\n\nRow 3 | Product: Gadget; Price: 19.50; In stock: no"


def test_csv_quoted_fields_and_cp1252():
    content = 'Name,Note\n"Rao, Alice","Paid 5€ fee"\n'.encode("cp1252")
    (section,) = load_document("people.csv", content)
    assert section.text == "Row 2 | Name: Rao, Alice; Note: Paid 5€ fee"


@pytest.mark.parametrize("name", ["broken.xlsx", "broken.xls"])
def test_corrupt_workbooks_are_rejected(name):
    with pytest.raises(DocumentParseError):
        load_document(name, b"this is not a spreadsheet")


def test_empty_csv_is_rejected():
    with pytest.raises(DocumentParseError):
        load_document("empty.csv", b",,,\n,,\n")


def test_spreadsheet_row_is_retrieved_and_cited(service, fake_llm):
    rows = [["Name", "Department", "Salary"]] + [[f"Person {i}", "Ops", 40000 + i] for i in range(60)]
    rows.insert(30, ["Bob Iyer", "Engineering", 61000])
    service.ingest("staff.xlsx", _xlsx({"Employees": rows}))
    answer = service.ask("What is Bob Iyer's salary?")
    assert answer.status == "answered"
    assert "Bob Iyer" in answer.sources[0].excerpt and "61000" in answer.sources[0].excerpt
    assert answer.sources[0].section == "Employees"


def test_upload_spreadsheets_through_api(client):
    files = [
        ("files", ("products.csv", b"Product,Price\nWidget,9.99\n")),
        ("files", ("inventory.xls", (FIXTURES / "inventory.xls").read_bytes())),
        ("files", ("offices.xlsx", _xlsx({"Offices": [["City", "Head count"], ["Chennai", 40]]}))),
    ]
    response = client.post("/documents", files=files)
    assert response.status_code == 201
    assert sorted(r["document"]["name"] for r in response.json()["results"]) == [
        "inventory.xls", "offices.xlsx", "products.csv",
    ]
