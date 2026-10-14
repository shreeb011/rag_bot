"""Generate large test documents to try uploads near the size limit.

    python scripts/make_large_files.py --size-mb 25                 # every format, ~25 MB each
    python scripts/make_large_files.py --size-mb 50 --formats csv pdf --out large_test_files

CSV and TXT files are all text, so their size is their text. Real PDF, Word and Excel files are mostly large
because of images (text compresses ~10x inside them), so those get ``--text-mb`` of text (default 1 MB,
already hundreds of pages) plus an embedded image that brings the file up to the requested size. That tests
the upload path at full size without turning one file into hundreds of thousands of chunks.

Every file contains one planted fact roughly in the middle, so you can check that retrieval still finds a
needle in a large document:

    "What is the access code for the Halden archive vault?"   ->  7319-KESTREL

Needs reportlab and Pillow for PDFs (both in requirements-dev.txt).
"""

from __future__ import annotations

import argparse
import csv
import io
import os
import random
from pathlib import Path

NEEDLE = (
    "Facilities notice: the access code for the Halden archive vault is 7319-KESTREL. "
    "Only the records manager may share it."
)
NEEDLE_QUESTION = "What is the access code for the Halden archive vault?"
MB = 1024 * 1024

_FIRST = "Asha Ravi Meena Karthik Priya Suresh Divya Arjun Lakshmi Vikram Nisha Rahul Kavya Anand Deepa".split()
_LAST = "Sharma Iyer Nair Raj Babu Menon Reddy Pillai Kumar Das Rao Shetty Varma Joshi Bose".split()
_DEPTS = "Finance Engineering Sales Marketing Operations Legal Support Procurement Research Facilities".split()
_CITIES = "Chennai Bengaluru Mumbai Pune Hyderabad Kochi Delhi Kolkata Coimbatore Mysuru".split()
_TOPICS = [
    "quarterly budget review", "vendor onboarding", "laptop refresh", "customer escalation", "data retention",
    "office relocation", "travel approval", "security training", "contract renewal", "release planning",
]
_FIELDS = ["Ticket", "Employee", "Department", "City", "Topic", "Amount (INR)", "Status", "Notes"]


def _record(rng: random.Random, i: int) -> list[str]:
    return [
        f"T-{100000 + i}",
        f"{rng.choice(_FIRST)} {rng.choice(_LAST)}",
        rng.choice(_DEPTS),
        rng.choice(_CITIES),
        rng.choice(_TOPICS),
        str(rng.randint(500, 900000)),
        rng.choice(["Open", "Closed", "Pending", "Escalated"]),
        "",
    ]


def _paragraphs(text_bytes: int) -> list[str]:
    rng = random.Random(7)
    out, size, i = [], 0, 0
    while size < text_bytes:
        ticket, who, dept, city, topic, amount, status, _ = _record(rng, i)
        text = (
            f"Ticket {ticket} was raised by {who} from the {dept} team in {city} about {topic}. The amount "
            f"involved is INR {amount} and the ticket is currently {status.lower()}. Follow-up notes were "
            f"added by the {rng.choice(_DEPTS)} desk."
        )
        out.append(text)
        size += len(text) + 2
        i += 1
    out.insert(len(out) // 2, NEEDLE)
    return out


def _rows(text_bytes: int) -> list[list[str]]:
    rng = random.Random(7)
    rows, size, i = [], 0, 0
    while size < text_bytes:
        row = _record(rng, i)
        rows.append(row)
        size += sum(len(c) for c in row) + len(row)
        i += 1
    rows.insert(len(rows) // 2, ["T-NOTICE", "", "", "", "", "", "", NEEDLE])
    return rows


def _noise_png(target_bytes: int) -> bytes:
    """An incompressible image of about ``target_bytes`` (random pixels don't compress)."""
    from PIL import Image

    if target_bytes <= 0:
        return b""
    side = max(8, int((target_bytes / 3) ** 0.5))
    image = Image.frombytes("RGB", (side, side), os.urandom(side * side * 3))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", compress_level=0)
    return buffer.getvalue()


def make_txt(path: Path, size_mb: float, text_mb: float) -> None:
    path.write_text("\n\n".join(_paragraphs(int(size_mb * MB))), encoding="utf-8")


def make_csv(path: Path, size_mb: float, text_mb: float) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(_FIELDS)
        writer.writerows(_rows(int(size_mb * MB)))


def make_xlsx(path: Path, size_mb: float, text_mb: float) -> None:
    import openpyxl
    from openpyxl.drawing.image import Image as SheetImage

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Tickets"
    sheet.append(_FIELDS)
    for row in _rows(int(text_mb * MB)):
        sheet.append(row)
    without_image = io.BytesIO()
    workbook.save(without_image)
    padding = int(size_mb * MB) - without_image.tell()
    if padding > 0:
        chart = workbook.create_sheet("Scanned form")
        chart.add_image(SheetImage(io.BytesIO(_noise_png(padding))), "A1")
    workbook.save(path)


def make_docx(path: Path, size_mb: float, text_mb: float) -> None:
    import docx
    from docx.shared import Inches

    document = docx.Document()
    document.add_heading("Service desk ticket log", 0)
    paragraphs = _paragraphs(int(text_mb * MB))
    for start in range(0, len(paragraphs), 400):
        document.add_heading(f"Batch {start // 400 + 1}", 1)
        for text in paragraphs[start : start + 400]:
            document.add_paragraph(text)
    without_image = io.BytesIO()
    document.save(without_image)
    padding = int(size_mb * MB) - without_image.tell()
    if padding > 0:
        document.add_heading("Appendix: scanned form", 1)
        document.add_picture(io.BytesIO(_noise_png(padding)), width=Inches(6))
    document.save(path)


def make_pdf(path: Path, size_mb: float, text_mb: float) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate

    style = getSampleStyleSheet()["BodyText"]
    texts = _paragraphs(int(text_mb * MB))
    without_image = io.BytesIO()
    SimpleDocTemplate(without_image, pagesize=A4, pageCompression=0).build([Paragraph(t, style) for t in texts])
    story = [Paragraph(t, style) for t in texts]  # flowables can't be laid out twice
    padding = int(size_mb * MB) - without_image.tell()
    if padding > 0:
        png = io.BytesIO(_noise_png(padding))
        width, height = ImageReader(png).getSize()
        png.seek(0)
        story += [PageBreak(), Image(png, width=400, height=400 * height / width)]
    SimpleDocTemplate(str(path), pagesize=A4, pageCompression=0).build(story)


MAKERS = {"txt": make_txt, "csv": make_csv, "xlsx": make_xlsx, "docx": make_docx, "pdf": make_pdf}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--size-mb", type=float, default=20, help="Target file size")
    parser.add_argument("--text-mb", type=float, default=1, help="Text inside PDF/DOCX/XLSX files")
    parser.add_argument("--formats", nargs="+", choices=sorted(MAKERS), default=sorted(MAKERS))
    parser.add_argument("--out", type=Path, default=Path("large_test_files"))
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    for fmt in args.formats:
        path = args.out / f"large_tickets.{fmt}"
        MAKERS[fmt](path, args.size_mb, min(args.text_mb, args.size_mb))
        print(f"{path}  {path.stat().st_size / MB:.1f} MB", flush=True)
    print(f"\nAsk: {NEEDLE_QUESTION!r}  (expected answer: 7319-KESTREL)")


if __name__ == "__main__":
    main()
