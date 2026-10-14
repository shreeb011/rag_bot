"""Turn raw uploaded bytes into a list of text ``Section``s, preserving page numbers and headings."""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from collections.abc import Iterable, Sequence
from pathlib import PurePath

import docx
import openpyxl
import xlrd
from docx.table import Table
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.errors import DocumentParseError, UnsupportedFileTypeError
from app.models import Section

SPREADSHEET_EXTENSIONS = {".csv", ".xlsx", ".xls"}
SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx"} | SPREADSHEET_EXTENSIONS

# Bump when parsing changes in a way that should re-index documents that were uploaded before.
PARSER_VERSION = 2

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_DOCX_HEADING_STYLE_RE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)
# "Email: a@b.com | Phone: +91 ..." -> one field per line; small LLMs misread long separator-joined lines.
_INLINE_FIELD_SEPARATOR_RE = re.compile(r"(?<=\S)[ \t]*[|\u2022\u00b7\u25aa\u2666\u25c6][ \t]*(?=[A-Z][A-Za-z/ ]{0,25}:\s)")


def file_extension(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def load_document(filename: str, data: bytes) -> list[Section]:
    ext = file_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        hint = " Save legacy .doc files as .docx first." if ext == ".doc" else ""
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext or filename}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}.{hint}"
        )
    if not data:
        raise DocumentParseError(f"'{filename}' is empty.")

    if ext == ".pdf":
        sections = _load_pdf(filename, data)
    elif ext == ".docx":
        sections = _load_docx(filename, data)
    elif ext == ".csv":
        sections = _load_csv(data)
    elif ext == ".xlsx":
        sections = _load_xlsx(filename, data)
    elif ext == ".xls":
        sections = _load_xls(filename, data)
    elif ext in {".md", ".markdown"}:
        sections = _load_markdown(_decode(data))
    else:
        sections = [Section(text=clean_text(_decode(data)))]

    sections = [s for s in sections if s.text.strip()]
    if not sections:
        raise DocumentParseError(f"No extractable text found in '{filename}'.")
    return sections


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join words hyphenated across lines
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = _INLINE_FIELD_SEPARATOR_RE.sub("\n", text)
    return text.strip()


def _load_pdf(filename: str, data: bytes) -> list[Section]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentParseError(f"'{filename}' is encrypted.")
        return [
            Section(text=_unwrap_lines(clean_text(page.extract_text() or "")), page=number)
            for number, page in enumerate(reader.pages, start=1)
        ]
    except DocumentParseError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise DocumentParseError(f"Could not parse PDF '{filename}': {exc}") from exc


def _unwrap_lines(text: str) -> str:
    """Re-join PDF lines that were hard-wrapped mid-sentence, keeping headings on their own line."""
    return re.sub(r"(?<=[,;])\n(?=\S)|(?<=\S)\n(?=[a-z(])|(?<=[^.!?:\s])\n(?=\d)", " ", text)


def _load_docx(filename: str, data: bytes) -> list[Section]:
    """Walk a Word document in reading order; headings become heading paths, tables become rows.

    Headings come from "Heading N" / "Title" styles. Many real documents (CVs especially) never use those
    styles, so when none are present, short ALL-CAPS or fully bold lines are treated as headings instead.
    """
    try:
        document = docx.Document(io.BytesIO(data))
        blocks = list(document.iter_inner_content())
    except Exception as exc:  # python-docx raises a mix of zip, XML and package errors for bad input
        raise DocumentParseError(f"Could not parse Word document '{filename}': {exc}") from exc

    uses_heading_styles = any(not isinstance(b, Table) and _docx_heading_level(b) is not None for b in blocks)
    sections: list[Section] = []
    path: list[tuple[int, str]] = []
    buffer: list[str] = []
    seen_body = False

    def flush() -> None:
        body = clean_text("\n".join(buffer))
        if body:
            heading = " > ".join(title for _, title in path) or None
            sections.append(Section(text=body, heading=heading))
        buffer.clear()

    for block in blocks:
        if isinstance(block, Table):
            buffer.extend(_docx_table_rows(block))
            buffer.append("")
            seen_body = True
            continue
        text = block.text.strip()
        if not text:
            continue
        level = _docx_heading_level(block)
        if level is None and not uses_heading_styles:
            level = _docx_implicit_heading_level(block, text)
            if level is not None and not seen_body:
                # Lines above the first paragraph (a name, a job title) are the document's header: keep them
                # as text so they stay searchable, and use the first one as the document title.
                buffer.append(text)
                path = path or [(0, text)]
                continue
        if level is not None:
            flush()
            path = [(lvl, t) for lvl, t in path if lvl < level] + [(level, text)]
            continue
        seen_body = True
        if _docx_style_name(block).lower().startswith("list"):
            buffer.append(f"- {text}")
        else:
            buffer.append(text)
            buffer.append("")
    flush()
    return sections


def _docx_style_name(paragraph) -> str:
    try:
        style = paragraph.style
    except (KeyError, ValueError):
        return ""
    return (style.name or "") if style is not None else ""


def _docx_heading_level(paragraph) -> int | None:
    name = _docx_style_name(paragraph)
    if name.lower() == "title":
        return 0
    match = _DOCX_HEADING_STYLE_RE.match(name)
    return int(match.group(1)) if match else None


def _docx_implicit_heading_level(paragraph, text: str) -> int | None:
    """1 for a short ALL-CAPS line, 2 for a short fully bold line, otherwise None."""
    if len(text) > 80 or text[-1] in ",;" or not any(c.isalpha() for c in text):
        return None
    if _docx_style_name(paragraph).lower().startswith("list"):
        return None
    if text == text.upper():  # may end in "." for initials, e.g. "JANE DOE M.S."
        return 1
    if text.endswith("."):
        return None
    try:
        style_bold = bool(paragraph.style is not None and paragraph.style.font.bold)
    except (KeyError, ValueError):
        style_bold = False
    runs = [r for r in paragraph.runs if r.text.strip()]
    if runs and all(r.bold or (r.bold is None and style_bold) for r in runs):
        return 2
    return None


def _docx_table_rows(table: Table) -> list[str]:
    rows = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            text = " ".join(cell.text.split())
            if text and (not cells or cells[-1] != text):  # merged cells repeat their text
                cells.append(text)
        if cells:
            row_text = cells[0]
            for cell in cells[1:]:
                row_text += f" {cell}" if row_text.endswith(":") else f" | {cell}"
            rows.append(row_text)
    return rows


def _load_csv(data: bytes) -> list[Section]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = csv.reader(io.StringIO(text), dialect)
    return _table_sections(None, ([cell.strip() for cell in row] for row in rows))


def _load_xlsx(filename: str, data: bytes) -> list[Section]:
    try:
        # data_only: formulas give the value Excel last calculated, not the formula text.
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises zip, XML and key errors for bad input
        raise DocumentParseError(f"Could not parse Excel workbook '{filename}': {exc}") from exc
    try:
        sections = []
        for sheet in workbook.worksheets:
            rows = ([_xlsx_cell_text(cell) for cell in row] for row in sheet.iter_rows(min_row=1, min_col=1))
            sections.extend(_table_sections(sheet.title, rows))
        return sections
    finally:
        workbook.close()


def _xlsx_cell_text(cell) -> str:
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in (getattr(cell, "number_format", "") or ""):
        return f"{value * 100:.10g}%"
    return _cell_text(value)


def _load_xls(filename: str, data: bytes) -> list[Section]:
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise DocumentParseError(f"Could not parse Excel workbook '{filename}': {exc}") from exc

    def value(cell) -> object:
        if cell.ctype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(cell.value, book.datemode)
            except (ValueError, OverflowError):
                return cell.value
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype == xlrd.XL_CELL_ERROR:
            return None
        return cell.value

    sections = []
    for sheet in book.sheets():
        rows = ([_cell_text(value(c)) for c in sheet.row(r)] for r in range(sheet.nrows))
        sections.extend(_table_sections(sheet.name, rows))
    return sections


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.10g}"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return " ".join(str(value).split())


def _table_sections(sheet: str | None, rows: Iterable[Sequence[str]]) -> list[Section]:
    """One section per sheet; each row becomes ``Row N | Header: value; Header: value``.

    Repeating the column names on every row keeps a chunk meaningful on its own (a chunk from the middle of
    a sheet doesn't contain the header row), which is what both retrieval and the answering LLM need. Row
    numbers match the spreadsheet's, so a cited row can be found in Excel.
    """
    lines: list[str] = []
    header: list[str] | None = None
    seen_data = False
    for number, row in enumerate(rows, start=1):
        cells = list(row)
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            continue
        filled = [cell for cell in cells if cell]
        if header is None and not seen_data:
            if _looks_like_header(cells):
                header = [cell or _column_letter(i) for i, cell in enumerate(cells)]
                continue
            if len(filled) == 1:  # a title above the table, e.g. "Employee list 2025"
                lines.append(filled[0])
                continue
        seen_data = True
        if header:
            fields = [
                f"{header[i] if i < len(header) else _column_letter(i)}: {cell}"
                for i, cell in enumerate(cells)
                if cell
            ]
        else:
            fields = filled
        lines.append(f"Row {number} | " + "; ".join(fields))
    if not lines:
        return []
    return [Section(text="\n\n".join(lines), heading=sheet)]


def _looks_like_header(cells: list[str]) -> bool:
    filled = [c for c in cells if c]
    return len(filled) >= 2 and all(not _is_number(c) for c in filled) and len(filled) == len(set(filled))


def _is_number(text: str) -> bool:
    try:
        float(text.replace(",", "").rstrip("%"))
        return True
    except ValueError:
        return False


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"Column {letters}"


def _load_markdown(text: str) -> list[Section]:
    """Split Markdown on headings; each section remembers its heading path (e.g. "Leave > Sick leave")."""
    sections: list[Section] = []
    path: list[tuple[int, str]] = []
    buffer: list[str] = []
    in_code_block = False

    def flush() -> None:
        body = clean_text("\n".join(buffer))
        if body:
            heading = " > ".join(title for _, title in path) or None
            sections.append(Section(text=body, heading=heading))
        buffer.clear()

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_code_block = not in_code_block
        match = None if in_code_block else _HEADING_RE.match(line)
        if match:
            flush()
            level, title = len(match.group(1)), match.group(2).strip()
            path = [(lvl, t) for lvl, t in path if lvl < level] + [(level, title)]
        else:
            buffer.append(line)
    flush()
    return sections
